#!/usr/bin/env python3
"""Generate synthetic QA fixtures for RegexFlow AI.

All values are synthetic. Card-like values are fake/test values only.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "tests" / "fixtures" / "generated"


def write_dataset(name: str, columns: list[str], rows: list[dict]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    write_csv(OUTPUT_DIR / f"{name}.csv", columns, rows)
    write_xlsx(OUTPUT_DIR / f"{name}.xlsx", columns, rows)


def write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    column: "" if row.get(column) is None else row.get(column)
                    for column in columns
                }
            )


def write_xlsx(path: Path, columns: list[str], rows: list[dict]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"

    for column_index, column in enumerate(columns, start=1):
        worksheet.cell(row=1, column=column_index, value=column)

    for row_index, row in enumerate(rows, start=2):
        for column_index, column in enumerate(columns, start=1):
            value = row.get(column)
            cell = worksheet.cell(row=row_index, column=column_index)
            if isinstance(value, str) and value.startswith("="):
                cell.value = value
                cell.data_type = "s"
            else:
                cell.value = value

    workbook.save(path)


def email_cases() -> tuple[list[str], list[dict]]:
    columns = ["ID", "Email", "ExpectedValid", "Notes"]
    rows = [
        {"ID": "E001", "Email": "john.doe@example.com", "ExpectedValid": "yes", "Notes": "standard lowercase email"},
        {"ID": "E002", "Email": "jane_smith@domain.com", "ExpectedValid": "yes", "Notes": "underscore local part"},
        {"ID": "E003", "Email": "user+tag@sub.example.org", "ExpectedValid": "yes", "Notes": "tag and subdomain"},
        {"ID": "E004", "Email": "UPPER.CASE@EXAMPLE.COM", "ExpectedValid": "yes", "Notes": "uppercase address"},
        {"ID": "E005", "Email": "invalid-email", "ExpectedValid": "no", "Notes": "missing domain"},
        {"ID": "E006", "Email": "missing_at_symbol.com", "ExpectedValid": "no", "Notes": "missing at sign"},
        {"ID": "E007", "Email": " leading@example.com ", "ExpectedValid": "yes", "Notes": "leading and trailing spaces"},
        {"ID": "E008", "Email": "", "ExpectedValid": "no", "Notes": "empty string"},
        {"ID": "E009", "Email": None, "ExpectedValid": "no", "Notes": "null cell"},
        {"ID": "E010", "Email": "john.doe@example.com", "ExpectedValid": "yes", "Notes": "duplicate value"},
    ]
    return columns, rows


def phone_cases() -> tuple[list[str], list[dict]]:
    columns = ["ID", "Phone", "ExpectedE164", "ExpectedValid", "Notes"]
    rows = [
        {"ID": "P001", "Phone": "0412 345 678", "ExpectedE164": "+61412345678", "ExpectedValid": "yes", "Notes": "AU mobile local"},
        {"ID": "P002", "Phone": "+61 412 345 678", "ExpectedE164": "+61412345678", "ExpectedValid": "yes", "Notes": "AU mobile international"},
        {"ID": "P003", "Phone": "03 9123 4567", "ExpectedE164": "+61391234567", "ExpectedValid": "yes", "Notes": "AU landline local"},
        {"ID": "P004", "Phone": "+61 3 9123 4567", "ExpectedE164": "+61391234567", "ExpectedValid": "yes", "Notes": "AU landline international"},
        {"ID": "P005", "Phone": "(03) 9123 4567", "ExpectedE164": "+61391234567", "ExpectedValid": "yes", "Notes": "parenthesized area code"},
        {"ID": "P006", "Phone": "12345", "ExpectedE164": "", "ExpectedValid": "no", "Notes": "too short"},
        {"ID": "P007", "Phone": "abc123", "ExpectedE164": "", "ExpectedValid": "no", "Notes": "letters and digits"},
        {"ID": "P008", "Phone": "", "ExpectedE164": "", "ExpectedValid": "no", "Notes": "empty string"},
        {"ID": "P009", "Phone": None, "ExpectedE164": "", "ExpectedValid": "no", "Notes": "null cell"},
    ]
    return columns, rows


def card_cases() -> tuple[list[str], list[dict]]:
    columns = ["ID", "CardNumber", "LuhnValid", "Notes"]
    rows = [
        {"ID": "C001", "CardNumber": "4111 1111 1111 1111", "LuhnValid": "yes", "Notes": "fake test Visa-like number"},
        {"ID": "C002", "CardNumber": "5555-5555-5555-4444", "LuhnValid": "yes", "Notes": "fake test Mastercard-like number"},
        {"ID": "C003", "CardNumber": "4000 0000 0000 0002", "LuhnValid": "yes", "Notes": "fake test card-like number"},
        {"ID": "C004", "CardNumber": "1234 5678 9012 3456", "LuhnValid": "no", "Notes": "Luhn-invalid sequence"},
        {"ID": "C005", "CardNumber": "0000 0000 0000 0000", "LuhnValid": "yes", "Notes": "synthetic all-zero Luhn-valid edge case"},
        {"ID": "C006", "CardNumber": "not-a-card", "LuhnValid": "no", "Notes": "plain text"},
        {"ID": "C007", "CardNumber": "", "LuhnValid": "no", "Notes": "empty string"},
        {"ID": "C008", "CardNumber": None, "LuhnValid": "no", "Notes": "null cell"},
    ]
    return columns, rows


def address_cases() -> tuple[list[str], list[dict]]:
    columns = ["ID", "Address", "ExpectedAddressLike", "Notes"]
    rows = [
        {"ID": "A001", "Address": "123 Collins St, Melbourne VIC 3000", "ExpectedAddressLike": "yes", "Notes": "street abbreviation"},
        {"ID": "A002", "Address": "Unit 5/88 Elizabeth Street, Melbourne VIC 3000", "ExpectedAddressLike": "yes", "Notes": "unit address"},
        {"ID": "A003", "Address": "42 George Road, Sydney NSW 2000", "ExpectedAddressLike": "yes", "Notes": "road address"},
        {"ID": "A004", "Address": "9 Queen Ave, Brisbane QLD 4000", "ExpectedAddressLike": "yes", "Notes": "avenue address"},
        {"ID": "A005", "Address": "No fixed address", "ExpectedAddressLike": "no", "Notes": "negative control"},
        {"ID": "A006", "Address": "invalid address text", "ExpectedAddressLike": "no", "Notes": "invalid text"},
        {"ID": "A007", "Address": "", "ExpectedAddressLike": "no", "Notes": "empty string"},
        {"ID": "A008", "Address": None, "ExpectedAddressLike": "no", "Notes": "null cell"},
    ]
    return columns, rows


def mixed_pii_dataset() -> tuple[list[str], list[dict]]:
    columns = [
        "ID",
        "Name",
        "Email",
        "Phone",
        "CardNumber",
        "Website",
        "Address",
        "Notes",
        "InvoiceId",
        "DateText",
    ]
    rows = [
        {
            "ID": "M001",
            "Name": "John Example",
            "Email": "john.doe@example.com",
            "Phone": "0412 345 678",
            "CardNumber": "4111 1111 1111 1111",
            "Website": "https://example.com",
            "Address": "123 Collins St, Melbourne VIC 3000",
            "Notes": "Contact john.doe@example.com or call 0412 345 678",
            "InvoiceId": "INV-1001",
            "DateText": "01/05/2026",
        },
        {
            "ID": "M002",
            "Name": "Jane Example",
            "Email": "jane_smith@domain.com",
            "Phone": "+61 412 345 678",
            "CardNumber": "5555-5555-5555-4444",
            "Website": "http://test.org/path",
            "Address": "Unit 5/88 Elizabeth Street, Melbourne VIC 3000",
            "Notes": "Website: https://example.com, card: 4111 1111 1111 1111",
            "InvoiceId": "INV-1002",
            "DateText": "02/05/2026",
        },
        {
            "ID": "M003",
            "Name": "Alex Synthetic",
            "Email": "user+tag@sub.example.org",
            "Phone": "03 9123 4567",
            "CardNumber": "4000 0000 0000 0002",
            "Website": "www.example.net",
            "Address": "42 George Road, Sydney NSW 2000",
            "Notes": "Address: 123 Collins St, Melbourne VIC 3000",
            "InvoiceId": "INV-1003",
            "DateText": "03/05/2026",
        },
        {
            "ID": "M004",
            "Name": "Upper Case",
            "Email": "UPPER.CASE@EXAMPLE.COM",
            "Phone": "+61 3 9123 4567",
            "CardNumber": "1234 5678 9012 3456",
            "Website": "https://sub.domain.com/path?x=1",
            "Address": "9 Queen Ave, Brisbane QLD 4000",
            "Notes": "No sensitive data here",
            "InvoiceId": "INV-1004",
            "DateText": "04/05/2026",
        },
        {
            "ID": "M005",
            "Name": "Invalid Email",
            "Email": "invalid-email",
            "Phone": "(03) 9123 4567",
            "CardNumber": "0000 0000 0000 0000",
            "Website": "not-a-url",
            "Address": "No fixed address",
            "Notes": "Bad phone abc123 and invalid-email",
            "InvoiceId": "BAD-1005",
            "DateText": "2026-05-05",
        },
        {
            "ID": "M006",
            "Name": "Missing At",
            "Email": "missing_at_symbol.com",
            "Phone": "12345",
            "CardNumber": "not-a-card",
            "Website": "",
            "Address": "invalid address text",
            "Notes": "Cards 1234 5678 9012 3456 and phone 12345 are invalid",
            "InvoiceId": "BAD-1006",
            "DateText": "05-06-2026",
        },
        {
            "ID": "M007",
            "Name": "Empty Values",
            "Email": "",
            "Phone": "",
            "CardNumber": "",
            "Website": "https://example.com/empty",
            "Address": "",
            "Notes": "",
            "InvoiceId": "INV-1007",
            "DateText": "07/05/2026",
        },
        {
            "ID": "M008",
            "Name": None,
            "Email": None,
            "Phone": None,
            "CardNumber": None,
            "Website": None,
            "Address": None,
            "Notes": None,
            "InvoiceId": None,
            "DateText": None,
        },
    ]
    return columns, rows


def edge_cases() -> tuple[list[str], list[dict]]:
    long_text = "This is a synthetic long cell. " * 80
    columns = [
        "ID",
        "Name",
        "Value",
        "Email",
        "Phone",
        "Website",
        "Notes",
        "FormulaLike",
        "NumericString",
    ]
    rows = [
        {"ID": "EDGE-001", "Name": "Zoë Example", "Value": "unicode name", "Email": "zoe@example.com", "Phone": "0412 345 678", "Website": "https://example.com", "Notes": "Unicode name row", "FormulaLike": "=1+1", "NumericString": "001234"},
        {"ID": "EDGE-002", "Name": "李雷 Test", "Value": "leading and trailing spaces", "Email": " spaced@example.com ", "Phone": " 03 9123 4567 ", "Website": " www.example.net ", "Notes": "  padded note  ", "FormulaLike": "'=1+1", "NumericString": "000000"},
        {"ID": "EDGE-003", "Name": "Duplicate", "Value": "DUPLICATE", "Email": "duplicate@example.com", "Phone": "+61 412 345 678", "Website": "https://example.com/dup", "Notes": "duplicate@example.com appears twice duplicate@example.com", "FormulaLike": "", "NumericString": "1234567890123456"},
        {"ID": "EDGE-004", "Name": "Duplicate", "Value": "DUPLICATE", "Email": "duplicate@example.com", "Phone": "+61 412 345 678", "Website": "https://example.com/dup", "Notes": "duplicate row", "FormulaLike": "", "NumericString": "1234567890123456"},
        {"ID": "EDGE-005", "Name": "Long Text", "Value": long_text, "Email": "", "Phone": "", "Website": "", "Notes": long_text + " Contact long.text@example.com", "FormulaLike": "", "NumericString": "999999"},
        {"ID": "EDGE-006", "Name": "Multi Email", "Value": "multiple emails", "Email": "first@example.com; second@example.org", "Phone": "", "Website": "", "Notes": "first@example.com and second@example.org", "FormulaLike": "", "NumericString": "42"},
        {"ID": "EDGE-007", "Name": "Multi Phone", "Value": "multiple phones", "Email": "", "Phone": "0412 345 678 / 03 9123 4567", "Website": "", "Notes": "Call 0412 345 678 or 03 9123 4567", "FormulaLike": "", "NumericString": "12345"},
        {"ID": "EDGE-008", "Name": "URL Email", "Value": "url and email", "Email": "url.email@example.com", "Phone": "", "Website": "https://example.com/contact?email=url.email@example.com", "Notes": "Email url.email@example.com with website https://example.com", "FormulaLike": "", "NumericString": "1000"},
        {"ID": "EDGE-009", "Name": "Card Phone", "Value": "card-like and phone-like", "Email": "", "Phone": "0412 345 678", "Website": "", "Notes": "Card 4111 1111 1111 1111 and phone 0412 345 678", "FormulaLike": "", "NumericString": "4111111111111111"},
        {"ID": "EDGE-010", "Name": "", "Value": "", "Email": "", "Phone": "", "Website": "", "Notes": "", "FormulaLike": "", "NumericString": ""},
        {"ID": "EDGE-011", "Name": None, "Value": None, "Email": None, "Phone": None, "Website": None, "Notes": None, "FormulaLike": None, "NumericString": None},
    ]
    return columns, rows


def medium_mixed_dataset() -> tuple[list[str], list[dict]]:
    columns, base_rows = mixed_pii_dataset()
    rows = []
    for index in range(600):
        source = dict(base_rows[index % len(base_rows)])
        source["ID"] = f"MM{index + 1:04d}"
        source["Name"] = f"Synthetic User {index + 1}"
        source["Email"] = (
            f"user{index + 1}@example.com"
            if index % 5 != 0
            else "invalid-email"
        )
        source["InvoiceId"] = f"INV-{2000 + index}"
        source["DateText"] = f"{(index % 28) + 1:02d}/05/2026"
        rows.append(source)

    return columns, rows


def main() -> None:
    datasets = {
        "mixed_pii_dataset": mixed_pii_dataset(),
        "email_cases": email_cases(),
        "phone_cases": phone_cases(),
        "card_cases": card_cases(),
        "address_cases": address_cases(),
        "edge_cases": edge_cases(),
        "medium_mixed_dataset": medium_mixed_dataset(),
    }
    for name, (columns, rows) in datasets.items():
        write_dataset(name, columns, rows)

    print(f"Generated {len(datasets) * 2} files under {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
